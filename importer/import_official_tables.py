#!/usr/bin/env python3
"""Importa a Tabela 8.1 oficial do Novo CAGED para todos os municípios do Brasil."""

import argparse
import hashlib
import os
import re
from datetime import date
from pathlib import Path

import openpyxl

RA_ARACATUBA = {
    '350110', '350210', '350280', '350420', '350440', '350510', '350620',
    '350640', '350650', '350770', '350775', '350810', '351100', '351190',
    '351250', '351650', '351680', '351690', '351710', '351780', '351820',
    '351890', '352044', '352300', '352650', '352725', '352770', '353010',
    '353210', '353286', '353320', '353330', '353730', '353740', '353770',
    '354440', '354805', '354840', '354925', '355230', '355255', '355520',
    '355630',
}

MONTHS = {
    'janeiro': 1,
    'fevereiro': 2,
    'março': 3,
    'abril': 4,
    'maio': 5,
    'junho': 6,
    'julho': 7,
    'agosto': 8,
    'setembro': 9,
    'outubro': 10,
    'novembro': 11,
    'dezembro': 12,
}


def integer(value):
    return 0 if value in (None, '', '-') else int(round(float(value)))


def competence_from_header(value):
    match = re.fullmatch(
        r'\s*([A-Za-zçÇãÃ]+)/(20\d{2})\s*',
        str(value or ''),
    )

    if not match:
        return None

    month = match.group(1).lower()

    if month not in MONTHS:
        return None

    return date(
        int(match.group(2)),
        MONTHS[month],
        1,
    )


def clean_municipality_name(value):
    name = str(value or '').strip()
    return re.sub(r'^[A-Z][a-z]-', '', name)


def municipality_code(value):
    """Aceita somente o código municipal de seis posições do Novo CAGED."""
    digits = re.sub(r'\D', '', str(value or '').split('.')[0])
    return digits[:6] if len(digits) >= 6 else ''


def valid_municipality_code(value):
    return bool(re.fullmatch(r'[1-5]\d{5}', value or ''))


def extract_records(workbook_path, source_url):
    workbook = openpyxl.load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    if 'Tabela 8.1' not in workbook.sheetnames:
        raise RuntimeError(
            'A planilha não possui a aba "Tabela 8.1".'
        )

    sheet = workbook['Tabela 8.1']

    headers = list(
        next(
            sheet.iter_rows(
                min_row=5,
                max_row=5,
                values_only=True,
            )
        )
    )

    month_columns = [
        (index, competence_from_header(value))
        for index, value in enumerate(headers)
        if competence_from_header(value)
    ]

    if not month_columns:
        raise RuntimeError(
            'Não foi possível localizar as competências da Tabela 8.1.'
        )

    records = []
    municipalities = {}
    excluded = []

    for row in sheet.iter_rows(min_row=7, values_only=True):
        if row[2] in (None, ''):
            continue

        ibge_code = municipality_code(row[2])
        name = clean_municipality_name(row[3])

        if not name:
            continue

        # Exclui “município não identificado” e linhas que não representam
        # municípios brasileiros. Esta é a causa da diferença dos totais.
        if not valid_municipality_code(ibge_code):
            excluded.append((ibge_code or str(row[2]), name))
            continue

        municipalities[ibge_code] = {
            'ibge_code': ibge_code,
            'name': name,
            'territory': (
                'Região Administrativa de Araçatuba'
                if ibge_code in RA_ARACATUBA
                else 'Brasil'
            ),
            'is_regional': ibge_code in RA_ARACATUBA,
        }

        for start, competence in month_columns:
            records.append({
                'competence': competence.isoformat(),
                'ibge_code': ibge_code,
                'stock': integer(row[start]),
                'admissions': integer(row[start + 1]),
                'dismissals': integer(row[start + 2]),
                'balance': integer(row[start + 3]),
                'source_url': source_url,
            })

    if len(municipalities) < 5000:
        raise RuntimeError(
            f'Foram localizados somente {len(municipalities)} municípios. '
            'A leitura da Tabela 8.1 falhou.'
        )

    if excluded:
        print(
            'Linhas geográficas não municipais excluídas: '
            + ', '.join(
                f'{code} ({name})'
                for code, name in excluded[:10]
            )
        )

    return records, list(municipalities.values())


def supabase_request(
    method,
    table,
    url,
    key,
    payload=None,
    query='',
):
    import requests

    headers = {
        'apikey': key,
        'Authorization': f'Bearer {key}',
        'Content-Type': 'application/json',
        'Prefer': 'resolution=merge-duplicates',
    }

    response = requests.request(
        method,
        f'{url}/rest/v1/{table}{query}',
        headers=headers,
        json=payload,
        timeout=180,
    )

    response.raise_for_status()


def save(records, municipalities, workbook_path, source_url):
    url = os.getenv('SUPABASE_URL')
    key = os.getenv('SUPABASE_SERVICE_ROLE_KEY')

    if not url or not key:
        raise RuntimeError(
            'Defina SUPABASE_URL e SUPABASE_SERVICE_ROLE_KEY.'
        )

    for start in range(0, len(municipalities), 500):
        supabase_request(
            'POST',
            'municipalities',
            url,
            key,
            municipalities[start:start + 500],
            '?on_conflict=ibge_code',
        )

    for start in range(0, len(records), 500):
        supabase_request(
            'POST',
            'caged_official_monthly',
            url,
            key,
            records[start:start + 500],
            '?on_conflict=competence,ibge_code',
        )

        if start and start % 25000 == 0:
            print(
                f'Importados {start} de {len(records)} '
                'registros mensais.'
            )

    dates = sorted(row['competence'] for row in records)

    metadata = {
        'source_url': source_url,
        'source_sha256': hashlib.sha256(
            workbook_path.read_bytes()
        ).hexdigest(),
        'competence_start': dates[0],
        'competence_end': dates[-1],
        'rows_imported': len(records),
    }

    supabase_request(
        'POST',
        'caged_official_imports',
        url,
        key,
        metadata,
        '?on_conflict=source_url',
    )

    supabase_request(
        'POST',
        'rpc/refresh_caged_official_national',
        url,
        key,
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--file', required=True)
    parser.add_argument('--source-url', required=True)
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()

    workbook_path = Path(args.file)

    if not workbook_path.is_file():
        parser.error(f'Arquivo não encontrado: {workbook_path}')

    records, municipalities = extract_records(
        workbook_path,
        args.source_url,
    )

    birigui = next(
        row
        for row in records
        if row['ibge_code'] == '350650'
        and row['competence'] == '2026-06-01'
    )

    print(
        f'Validação Birigui 2026-06: '
        f'estoque {birigui["stock"]}; '
        f'admissões {birigui["admissions"]}; '
        f'desligamentos {birigui["dismissals"]}; '
        f'saldo {birigui["balance"]}.'
    )

    print(
        f'Municípios localizados: {len(municipalities)}. '
        f'Registros mensais: {len(records)}.'
    )

    if args.dry_run:
        print('Validação concluída sem enviar dados ao Supabase.')
    else:
        save(
            records,
            municipalities,
            workbook_path,
            args.source_url,
        )

        print(
            f'Importação nacional concluída: {len(records)} registros, '
            f'de {records[0]["competence"]} '
            f'a {records[-1]["competence"]}.'
        )


if __name__ == '__main__':
    main()
