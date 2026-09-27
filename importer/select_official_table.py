#!/usr/bin/env python3
"""Escolhe a planilha Tabela 8.1 com a competência mais recente."""

import argparse
import re
from datetime import date
from pathlib import Path

import openpyxl


MONTHS = {
    "janeiro": 1, "fevereiro": 2, "março": 3, "abril": 4,
    "maio": 5, "junho": 6, "julho": 7, "agosto": 8,
    "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}


def latest_competence(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Tabela 8.1" not in workbook.sheetnames:
        return None
    headers = next(workbook["Tabela 8.1"].iter_rows(min_row=5, max_row=5, values_only=True))
    competences = []
    for value in headers:
        match = re.fullmatch(r"\s*([A-Za-zçÇãÃ]+)/(20\d{2})\s*", str(value or ""))
        if match and match.group(1).lower() in MONTHS:
            competences.append(date(int(match.group(2)), MONTHS[match.group(1).lower()], 1))
    return max(competences, default=None)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--folder", required=True)
    parser.add_argument("--competence", help="AAAAMM que a planilha precisa conter.")
    args = parser.parse_args()
    candidates = []
    for path in Path(args.folder).rglob("*.xlsx"):
        competence = latest_competence(path)
        if competence:
            candidates.append((competence, path))
    if not candidates:
        raise SystemExit("Nenhuma planilha com a aba Tabela 8.1 foi encontrada.")
    latest, path = max(candidates, key=lambda item: item[0])
    if args.competence:
        match = re.fullmatch(r"(20\d{2})(0[1-9]|1[0-2])", args.competence)
        if not match:
            raise SystemExit("Use --competence no formato AAAAMM.")
        requested = date(int(match.group(1)), int(match.group(2)), 1)
        if latest < requested:
            raise SystemExit(
                f"A Tabela 8.1 mais recente é {latest:%Y%m}, mas a competência "
                f"solicitada é {requested:%Y%m}. Envie a planilha revisada ao Google Drive."
            )
    print(path)


if __name__ == "__main__":
    main()
